"""Export approved documents to JSON, CSV or Excel (one sheet per document type, plus the audit trail)."""
from __future__ import annotations

import csv
import io
import json
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font

from .normalize import DECIMALS
from .schemas import SCHEMAS


def _row(d: dict) -> dict:
    out = dict(id=d["id"], filename=d["filename"], document_type=d["document_type"], status=d["status"],
               extractor=d["extractor"], engine=d["engine"], reviewed_by=d.get("reviewer"))
    for name, f in (d["fields"] or {}).items():
        out[name] = f["value"]
        out[f"{name}_confidence"] = f["confidence"]
    return out


def to_json(docs: list[dict]) -> bytes:
    body = [{k: d[k] for k in ("id", "filename", "document_type", "status", "extractor", "engine", "reviewer", "fields")} for d in docs]
    return json.dumps(body, ensure_ascii=False, indent=2).encode()


def to_csv(docs: list[dict]) -> bytes:
    rows = [_row(d) for d in docs]
    cols: list[str] = []
    for r in rows:
        cols += [k for k in r if k not in cols]
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=cols)
    w.writeheader()
    w.writerows(rows)
    return ("﻿" + buf.getvalue()).encode()  # BOM so Excel opens Arabic correctly


def to_xlsx(docs: list[dict], audit: list[dict]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sid, schema in SCHEMAS.items():
        subset = [d for d in docs if d["document_type"] == sid]
        if not subset:
            continue
        ws = wb.create_sheet(sid[:31])
        head = ["id", "filename", "status", "reviewed_by"] + [x for f in schema.fields for x in (f.name, f.name + "_confidence")]
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
        for d in subset:
            r = _row(d)
            row = []
            for h in head:
                v = r.get(h)
                f = next((f for f in schema.fields if f.name == h), None)
                if f and f.type in ("money", "percent") and v is not None:
                    v = float(Decimal(v))
                row.append(v)
            ws.append(row)
            for f in schema.fields:
                if f.type == "money":
                    cur = (d["fields"].get("currency") or {}).get("value")
                    ws.cell(ws.max_row, head.index(f.name) + 1).number_format = "#,##0." + "0" * DECIMALS.get(cur, 2)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(40, max(10, max(len(str(c.value or "")) for c in col) + 2))
    a = wb.create_sheet("audit_trail")
    a.append(["time (UTC)", "document", "actor", "action", "field", "old", "new", "note"])
    for e in audit:
        a.append([e["ts"], e["filename"], e["actor"], e["action"], e["field"], e["old"], e["new"], e["note"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
